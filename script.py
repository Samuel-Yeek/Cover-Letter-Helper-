import pandas as pd
from docx import Document
import os
import win32com.client as win32

def generate_and_email_documents(excel_file, output_folder, template):
    # Read the Excel file
    try:
        data = pd.read_excel(excel_file)
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        return
    
    # Ensure the output folder exists
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    # Create an instance of Outlook
    outlook = win32.Dispatch('outlook.application')

    # Iterate through each row in the Excel file
    for index, row in data.iterrows():
        # Create a new document based on the template
        doc = Document(template)

        # Replace placeholders in the template with Excel data
        for paragraph in doc.paragraphs:
            for key in data.columns:
                placeholder = f"{{{{{key}}}}}"  # Placeholder format {{column_name}}
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, str(row[key]))

        # Save the Word document
        file_name = f"Document_{index + 1}.docx"
        file_path = os.path.join(output_folder, file_name)
        doc.save(file_path)
        print(f"Generated: {file_name}")

        # Prepare the email
        mail = outlook.CreateItem(0)  # 0 = MailItem
        mail.Subject = f"Application for {row['Position Name']} at {row['Company Name']}"
        mail.To = row.get('Recipient Email', '') 
        mail.Body = "Please find my application attached."

        # Convert Word content to email body
        word = win32.Dispatch('Word.Application')
        doc = word.Documents.Open(file_path)
        doc.Content.Copy()  # Copy the content of the Word document
        mail.Body = doc.Content.Text  # Paste it into the email body
        doc.Close(False)

        # Attach the Word document
        mail.Attachments.Add(file_path)

        # Send the email
        try:
            mail.Send()
            print(f"Email sent to {row.get('Recipient Email', 'No Email Provided')} for {file_name}")
        except Exception as e:
            print(f"Failed to send email for {file_name}: {e}")

# Example usage
excel_file = "cover_letter_data.xlsx"  # Path to your Excel file
output_folder = "output_documents"    # Folder to save generated documents
template = "template.docx"            # Path to your Word template with placeholders

generate_and_email_documents(excel_file, output_folder, template)
